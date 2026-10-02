# -*- coding: utf-8 -*-
"""把 md 5.5 节插入 结题报告.docx（python-docx，纯文本段落 + 局部加粗 run）。

锚点：正文标题「6. 工程与可靠性」（取最后一次出现，避开目录条目），
在其前 insert_paragraph_before 插入 5.5 全部内容。
DRY=1 时只打印锚点/样式诊断，不写文件。
"""
import os
import sys
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(errors="replace")
DOCX = Path(r"E:\707\docs\结题报告.docx")
DRY = os.environ.get("DRY") == "1"

doc = Document(str(DOCX))

# --- 定位锚点与样式参照 -------------------------------------------------------
anchor = None
h54 = None
paras = list(doc.paragraphs)
for i, p in enumerate(paras):
    t = p.text.strip()
    if t == "6. 工程与可靠性":
        anchor = p            # 逐次覆盖 → 最后一次出现 = 正文标题
    if t == "5.4 天花板声明":
        h54 = p
        h54_idx = i

assert anchor is not None, "未找到「6. 工程与可靠性」段落"
print("锚点段落 style:", anchor.style.name, "| 文本:", anchor.text[:30])
print("5.4 标题 style:", h54.style.name if h54 else "未找到")
# 5.4 正文尾段（「本轨道交付的是…」）拿正文样式
body_ref = None
if h54 is not None:
    for p in paras[h54_idx + 1: h54_idx + 6]:
        if p.text.strip().startswith("本轨道交付"):
            body_ref = p
            break
print("5.4 正文段 style:", body_ref.style.name if body_ref else "未找到")

TITLE = "5.5 策略升级：把 1.484× 变成可执行模块（2026-09-25）"
LEAD = ("前几节给出结论“冷门组合中奖时实得 1.484×”；本节把该结论拆成六个可计算、可检验的模块。"
        "所有模块只作用于奖金分摊侧，不触碰概率侧——选号不改中奖率（期望线性性）这一边界在任何模块中都不放宽。")
SECTIONS = [
    ("① 三口径 EV 对照（ev/total_ev.py）：",
     "同一注号码，同时给出三种分摊口径的单注 EV——名义（官方赔率表）、λ 均摊（随机注平均同注数）、"
     "冷/热门口径（流行度模型给出的 μ_q）。固定奖级在三种口径下逐位相等（单测锁死），差异只出现在浮动奖级"
     "——这就是“冷热差别全在分摊”的可视化。对照表附反事实回测的 1.484× 作为历史参照。"),
    ("② 冷门口径前置（ev/rollover.py / ev/sales_model.py）：",
     "rollover 的 λ（随机注平均同注数）升级为可选的 μ_冷门 = λ·exp(β·f(q10 组合))。冷门口径下参与人数阈值更低、"
     "正 EV 窗口的判定更贴近“买冷门”的真实决策；原 λ 键全部保留（看板 honest 页依赖），新键并存并标注口径。"
     "模型 NO-Go 时显式降级并说明原因，绝不静默回退。"),
    ("③ 奖池参与门槛（ev/threshold.py）：",
     "给定“冷门组合 + 分摊占比 0.7 + 顶格判定 θ=0.999”，解析解与数值解（二分 pool_share_expected）双口径给出"
     "奖池需达到多少才值得参与。实测结论：冷门口径门槛显著低于随机口径——这就是 1.484× 的门槛翻译。"
     "报告同时核对资金约束与注数上限，给出三档判定。"),
    ("④ Kelly 纪律仓位（ev/stake.py）：",
     "EV≤0 ⇔ f*≤0 → 建议注数 0（“纪律不下注”）。EV 为正时按 Kelly 公式给出理论注数，再与预算、单期注数上限、"
     "止损/熔断状态取最小——信号再好也不会击穿资金纪律。四分之一 Kelly 优先于半 Kelly（估计误差保护）。"),
    ("⑤ 限号前置（ev/selection.py）：",
     "七星彩的排序键本就是限号EV（crowd v2 撞号分薄），本次把每注的限号暴露概率与同号期望注数 μ 前置存储并回填到"
     "选中行，返回标注“已前置（排序键=限号EV）”；双/大为浮动奖池制，明确标注“不适用”；固定赔率彩种维持拒绝并指引"
     " --sample 反大众形态采样。"),
    ("⑥ 组合方差指标（ev/selection.py）：",
     "选中注集输出平均/最大重叠与有效独立注数 m_eff = m²/ΣΣoᵢⱼ（按主区号码归一，自重叠=1）。注注相同 → m_eff=1"
     "（方差无分散），零重叠 → m_eff=m。低重叠只降低中奖结果之间的相关性（方差），不提高总期望"
     "——指标口径里写明“只降方差不提EV”。"),
]

if DRY:
    print("DRY 模式：不写文件。将插入 1 标题 + 1 引言 + %d 小节段。" % len(SECTIONS))
    sys.exit(0)

h_style = h54.style if h54 is not None else anchor.style
b_style = body_ref.style if body_ref is not None else anchor.style


def insert_para(before_p, text, style, bold_prefix=None):
    new = before_p.insert_paragraph_before(text="", style=style)
    if bold_prefix:
        r1 = new.add_run(bold_prefix)
        r1.bold = True
        new.add_run(text)
    else:
        new.add_run(text)
    return new


insert_para(anchor, TITLE, h_style)
insert_para(anchor, LEAD, b_style)
for prefix, rest in SECTIONS:
    insert_para(anchor, rest, b_style, bold_prefix=prefix)

doc.save(str(DOCX))
print("已插入并保存。")

# --- 复检：重开文档验证 -------------------------------------------------------
doc2 = Document(str(DOCX))
texts = [p.text.strip() for p in doc2.paragraphs]
assert TITLE in texts, "复检失败：5.5 标题缺失"
n_sec = sum(1 for t in texts if t.startswith(("① ", "② ", "③ ", "④ ", "⑤ ", "⑥ ")))
i55 = texts.index(TITLE)
i6 = len(texts) - 1 - texts[::-1].index("6. 工程与可靠性")
assert i55 < i6, "复检失败：5.5 不在 6 章之前"
print("复检 OK：5.5 标题在位，小节段 %d 个，位于 6 章之前" % n_sec)
